"""Monta a planilha de controle do gerador de cartas."""

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path("/home/claude/hr-synthetic-data-br/data/synthetic")
SAIDA = Path("/home/claude/gerador-cartas/Gerador_Cartas.xlsx")

FONTE = "Arial"
AZUL = "1F3864"
AZUL_INPUT = "0000FF"
AMARELO = "FFF2CC"
MOEDA = 'R$ #,##0.00'
PCT = "0.0%"

wb = Workbook()
wb.remove(wb.active)


def cabecalho(ws, linha=1, n=None):
    n = n or ws.max_column
    for c in range(1, n + 1):
        cel = ws.cell(row=linha, column=c)
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        cel.fill = PatternFill("solid", fgColor=AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[linha].height = 28
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def larguras(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


# =========================================================================
# Leia-me
# =========================================================================
ws = wb.create_sheet("Leia-me")
texto = [
    ("GERADOR DE CARTAS E COMUNICADOS", "t"),
    ("", None),
    ("O que faz", "h"),
    ("Preenche um modelo do Word para cada linha da aba Dados, salva em PDF "
     "individual, registra no log e opcionalmente prepara os e-mails no Outlook.", None),
    ("", None),
    ("Ordem de uso", "h"),
    ("1. Preencha a aba Config. Só as células azuis.", None),
    ("2. Preencha ou cole a lista na aba Dados.", None),
    ("3. Alt+F8 > ValidarDados. Corrija tudo que aparecer como Erro.", None),
    ("4. Deixe MODO_TESTE em SIM na primeira vez. Gera só as 3 primeiras cartas.", None),
    ("5. Alt+F8 > GerarCartas. Confira os PDFs.", None),
    ("6. Troque MODO_TESTE para NAO e gere o lote completo.", None),
    ("7. Alt+F8 > PrepararEmails, se for enviar por e-mail.", None),
    ("", None),
    ("Sobre os marcadores", "h"),
    ("Cada coluna da aba Dados vira um marcador no Word, em maiúsculas e entre "
     "chaves duplas. A coluna 'nome' vira {{NOME}}, 'salario_novo' vira "
     "{{SALARIO_NOVO}}.", None),
    ("Os parâmetros da aba Config viram marcadores do mesmo jeito: {{EMPRESA}}, "
     "{{CIDADE}}, {{ASSINATURA_NOME}}.", None),
    ("Acrescentar uma coluna na aba Dados cria um marcador novo sem mexer em "
     "código. É assim que o gerador serve para carta de promoção, de "
     "desligamento ou de qualquer outra coisa.", None),
    ("", None),
    ("Marcadores automáticos", "h"),
    ("{{DATA_CARTA}}      data por extenso, para o corpo da carta", None),
    ("{{DATA_GERACAO}}    data e hora da geração, para o rodapé", None),
    ("{{REFERENCIA}}      código único da carta, usado no log e na auditoria", None),
    ("", None),
    ("Por que os e-mails não são enviados automaticamente", "h"),
    ("PrepararEmails cria os rascunhos no Outlook e para. Quem aperta enviar é "
     "você. Envio automático de comunicado de remuneração é o tipo de "
     "automação que só dá errado uma vez.", None),
    ("", None),
    ("Legenda", "h"),
    ("Azul: célula de entrada, pode editar.", "i"),
    ("Fundo amarelo: revisar a cada ciclo.", "a"),
]
for i, (t, tipo) in enumerate(texto, start=1):
    cel = ws.cell(row=i, column=1, value=t)
    if tipo == "t":
        cel.font = Font(name=FONTE, bold=True, size=16, color=AZUL)
    elif tipo == "h":
        cel.font = Font(name=FONTE, bold=True, size=11, color=AZUL)
    elif tipo == "i":
        cel.font = Font(name=FONTE, size=10, color=AZUL_INPUT)
    elif tipo == "a":
        cel.font = Font(name=FONTE, size=10)
        cel.fill = PatternFill("solid", fgColor=AMARELO)
    else:
        cel.font = Font(name=FONTE, size=10)
larguras(ws, {"A": 98})

# =========================================================================
# Config
# =========================================================================
ws = wb.create_sheet("Config")
params = [
    ("Parâmetro", "Valor", "Nome definido", "Observação"),
    ("Pasta dos modelos", r"C:\Cartas\modelos", "PASTA_MODELOS",
     "Onde ficam os arquivos .docx"),
    ("Pasta de saída", r"C:\Cartas\saida", "PASTA_SAIDA",
     "Onde os PDFs serão gravados. Criada se não existir."),
    ("Nome da empresa", "Empresa Exemplo S.A.", "EMPRESA", "Marcador {{EMPRESA}}"),
    ("Cidade", "São Paulo", "CIDADE", "Marcador {{CIDADE}}"),
    ("Ciclo", "2026", "CICLO", "Marcador {{CICLO}}"),
    ("Assinatura - nome", "preencher", "ASSINATURA_NOME", "Marcador {{ASSINATURA_NOME}}"),
    ("Assinatura - cargo", "Diretoria de Pessoas", "ASSINATURA_CARGO",
     "Marcador {{ASSINATURA_CARGO}}"),
    ("Modo teste", "SIM", "MODO_TESTE",
     "SIM gera apenas as primeiras linhas. Deixe em SIM na primeira rodada."),
    ("Quantidade no modo teste", 3, "QTD_TESTE", "Quantas cartas gerar em modo teste"),
    ("Gerar PDF", "SIM", "GERAR_PDF", "NAO deixa o .docx preenchido em vez do PDF"),
    ("Limite de aumento para alerta", 0.30, "LIMITE_PCT",
     "REVISAR A CADA CICLO. Percentual acima do qual a validação emite aviso."),
    ("Assunto do e-mail", "Comunicado de mérito - ciclo {{CICLO}}", "ASSUNTO_EMAIL",
     "Aceita marcadores"),
    ("Corpo do e-mail", "{{TRATAMENTO}} {{NOME}}, segue em anexo seu comunicado.",
     "CORPO_EMAIL", "Aceita marcadores"),
]
for linha in params:
    ws.append(list(linha))
cabecalho(ws, 1, 4)

fmt = {10: "0", 12: PCT}
for r in range(2, len(params) + 1):
    for c in range(1, 5):
        ws.cell(row=r, column=c).font = Font(name=FONTE, size=10)
    cel = ws.cell(row=r, column=2)
    cel.font = Font(name=FONTE, size=10, color=AZUL_INPUT, bold=True)
    if r in fmt:
        cel.number_format = fmt[r]
    ws.cell(row=r, column=3).font = Font(name=FONTE, size=9, italic=True)
    ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")

ws.cell(row=12, column=2).fill = PatternFill("solid", fgColor=AMARELO)
dv = DataValidation(type="list", formula1='"SIM,NAO"', allow_blank=False)
ws.add_data_validation(dv)
dv.add("B9")
dv.add("B11")
larguras(ws, {"A": 32, "B": 42, "C": 22, "D": 54})

for r in range(2, len(params) + 1):
    nome = ws.cell(row=r, column=3).value
    wb.defined_names.add(DefinedName(nome, attr_text=f"Config!$B${r}"))

# =========================================================================
# Modelos
# =========================================================================
ws = wb.create_sheet("Modelos")
modelos = pd.DataFrame([
    ("MERITO", "Modelo_Merito.docx", "Comunicado de ajuste por mérito",
     "Carta_Merito_{{MATRICULA}}"),
    ("PROMOCAO", "Modelo_Promocao.docx", "Comunicado de promoção",
     "Carta_Promocao_{{MATRICULA}}"),
], columns=["cod_modelo", "arquivo", "descricao", "padrao_nome_arquivo"])
ws.append(list(modelos.columns))
for r in modelos.itertuples(index=False):
    ws.append(list(r))
cabecalho(ws, 1, 4)
for r in range(2, len(modelos) + 2):
    for c in range(1, 5):
        ws.cell(row=r, column=c).font = Font(name=FONTE, size=10)
larguras(ws, {"A": 14, "B": 26, "C": 34, "D": 34})
ws.cell(row=len(modelos) + 3, column=1,
        value="Para criar um modelo novo: copie um .docx existente, troque o texto, "
              "mantenha os marcadores e cadastre a linha aqui.").font = Font(
    name=FONTE, size=9, italic=True, color="808080")

# =========================================================================
# Dados (exemplo a partir da base sintética)
# =========================================================================
rng = np.random.default_rng(11)
snaps = pd.read_parquet(BASE / "fato_headcount_mensal.parquet")
colab = pd.read_parquet(BASE / "dim_colaborador.parquet")
areas = pd.read_parquet(BASE / "dim_area.parquet")
cargos = pd.read_parquet(BASE / "dim_cargo.parquet")

ult = snaps[snaps["id_mes"] == snaps["id_mes"].max()]
amostra = ult[ult["performance"] >= 4].sample(40, random_state=11).copy()
amostra = amostra.merge(colab[["matricula", "nome", "genero"]], on="matricula") \
                 .merge(areas[["id_area", "gerencia"]], on="id_area") \
                 .merge(cargos[["id_cargo", "cargo", "nivel"]], on="id_cargo")

prox = cargos.set_index(["familia_cargo", "grade"])["cargo"].to_dict()
amostra["cargo_novo"] = [
    prox.get((f, g + 1), c) for f, g, c in
    zip(amostra["familia_cargo"], amostra["grade"], amostra["cargo"])
]
amostra["pct"] = np.where(amostra["performance"] >= 5, 0.07, 0.045)
amostra["modelo"] = np.where(rng.random(len(amostra)) < 0.25, "PROMOCAO", "MERITO")
amostra.loc[amostra["modelo"] == "PROMOCAO", "pct"] = 0.12

dados = pd.DataFrame({
    "matricula": amostra["matricula"].astype(int),
    "nome": amostra["nome"],
    "email": amostra["nome"].str.lower().str.normalize("NFKD")
        .str.encode("ascii", "ignore").str.decode("ascii")
        .str.replace(" ", ".", regex=False) + "@empresaexemplo.com.br",
    "tratamento": np.where(amostra["genero"] == "Feminino", "Prezada", "Prezado"),
    "cargo_atual": amostra["cargo"],
    "cargo_novo": amostra["cargo_novo"],
    "gerencia": amostra["gerencia"],
    "gestor": "a definir",
    "salario_atual": amostra["salario"].round(2),
    "salario_novo": (amostra["salario"] * (1 + amostra["pct"])).round(2),
    "pct_aumento": amostra["pct"].round(4),
    "data_efeito": pd.Timestamp("2026-04-01").date(),
    "modelo": amostra["modelo"],
    "paragrafo_final": "Seguimos à disposição para esclarecer qualquer dúvida "
                       "por meio do seu business partner de RH.",
})
dados = dados.sort_values("matricula").reset_index(drop=True)

ws = wb.create_sheet("Dados")
ws.append(list(dados.columns))
for r in dados.itertuples(index=False):
    ws.append(list(r))
cabecalho(ws, 1, len(dados.columns))
formatos = {"salario_atual": MOEDA, "salario_novo": MOEDA,
            "pct_aumento": PCT, "data_efeito": "DD/MM/AAAA"}
for c, col in enumerate(dados.columns, start=1):
    for r in range(2, len(dados) + 2):
        cel = ws.cell(row=r, column=c)
        cel.font = Font(name=FONTE, size=10)
        if col in formatos:
            cel.number_format = formatos[col]
ref = f"A1:{get_column_letter(len(dados.columns))}{len(dados) + 1}"
t = Table(displayName="tblDados", ref=ref)
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws.add_table(t)
larguras(ws, {"A": 11, "B": 24, "C": 34, "D": 12, "E": 26, "F": 26, "G": 22,
              "H": 16, "I": 15, "J": 15, "K": 13, "L": 13, "M": 12, "N": 52})

# =========================================================================
# Validacao e Log
# =========================================================================
ws = wb.create_sheet("Validacao")
ws.append(["linha", "matricula", "severidade", "campo", "mensagem"])
cabecalho(ws, 1, 5)
larguras(ws, {"A": 9, "B": 12, "C": 13, "D": 20, "E": 78})

ws = wb.create_sheet("Log")
ws.append(["data_hora", "referencia", "matricula", "nome", "modelo",
           "arquivo_gerado", "verificacao", "status", "mensagem"])
cabecalho(ws, 1, 9)
larguras(ws, {"A": 20, "B": 26, "C": 12, "D": 24, "E": 12,
              "F": 40, "G": 14, "H": 12, "I": 44})

wb._sheets = [wb[n] for n in ["Leia-me", "Config", "Dados", "Modelos",
                              "Validacao", "Log"]]
SAIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SAIDA)
print("Salvo:", SAIDA, "| linhas de exemplo:", len(dados))
